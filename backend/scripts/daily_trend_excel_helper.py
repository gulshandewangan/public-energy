import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADERS = [
    "date",
    "timestamp",
    "demand",
    "frequency",
    "state_ui",
    "thermal",
    "thermal_ipp",
    "hydro",
    "wind",
    "solar",
    "total_generation",
    "renewable_generation",
    "renewable_share_pct",
    "pavagada_kspdcl",
    "bescom",
    "hescom",
    "gescom",
    "cesc",
    "mescom",
]
SHEET_NAME = "Daily Trends"


def ensure_workbook(path: Path):
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        sheet.title = SHEET_NAME
        if sheet.max_row == 1 and sheet["A1"].value is None:
            sheet.append(HEADERS)
        return workbook, sheet

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    return workbook, sheet


def sheet_to_records(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(value) for value in rows[0]]
    records = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        record = {}
        for index, header in enumerate(headers):
            record[header] = row[index] if index < len(row) else None
        records.append(record)
    return records


def upsert(path: Path):
    payload = json.loads(sys.stdin.read())
    workbook, sheet = ensure_workbook(path)
    records = sheet_to_records(sheet)

    row_number = None
    for index, record in enumerate(records, start=2):
        if str(record.get("date")) == str(payload["date"]):
            row_number = index
            break

    if row_number is None:
        row_number = sheet.max_row + 1

    for column_index, header in enumerate(HEADERS, start=1):
        sheet.cell(row=row_number, column=column_index, value=payload.get(header))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    print(json.dumps({"status": "ok", "row": row_number}))


def read(path: Path):
    if not path.exists():
        print("[]")
        return

    workbook, sheet = ensure_workbook(path)
    records = sheet_to_records(sheet)
    records.sort(key=lambda item: item["date"])
    print(json.dumps(records))


def main():
    command = sys.argv[1]
    workbook_path = Path(sys.argv[2]).resolve()

    if command == "upsert":
        upsert(workbook_path)
        return

    if command == "read":
        read(workbook_path)
        return

    raise SystemExit(f"Unsupported command: {command}")


if __name__ == "__main__":
    main()
